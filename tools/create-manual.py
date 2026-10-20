#!/usr/bin/env python3
"""Create Vietnamese v1.2 manual from one content file; render with render_docx.py."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1]
M=json.loads((R/'docs/manual-v1.2.json').read_text())
O=R/'dist/manual';O.mkdir(parents=True,exist_ok=True)
D=Document();S=D.sections[0];S.page_width=Inches(8.27);S.page_height=Inches(11.69)
S.top_margin=Inches(.62);S.bottom_margin=Inches(.62);S.left_margin=Inches(.72);S.right_margin=Inches(.72)
S.header_distance=Inches(.27);S.footer_distance=Inches(.3)
styles=D.styles
# Remove decorative paragraph rules inherited from the document template.
for style in styles:
 if style._element.pPr is not None:
  for border in list(style._element.pPr.findall(qn('w:pBdr'))):
   style._element.pPr.remove(border)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption','List Bullet','List Number']:
 st=styles[name];st.font.name='DejaVu Sans';st.font.color.rgb=RGBColor.from_string('222C28')
 st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'),'DejaVu Sans')
 st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'),'DejaVu Sans')
styles['Normal'].font.size=Pt(10.25);styles['Normal'].paragraph_format.line_spacing=1.08
styles['Normal'].paragraph_format.space_after=Pt(6)
styles['Title'].font.size=Pt(30);styles['Title'].font.bold=True;styles['Title'].font.color.rgb=RGBColor(0,0,0)
styles['Heading 1'].font.size=Pt(20);styles['Heading 1'].font.bold=True;styles['Heading 1'].paragraph_format.space_after=Pt(11)
styles['Heading 2'].font.size=Pt(11.5);styles['Heading 2'].font.bold=True;styles['Heading 2'].font.color.rgb=RGBColor.from_string('145D45');styles['Heading 2'].paragraph_format.space_before=Pt(10);styles['Heading 2'].paragraph_format.space_after=Pt(5)
styles['Caption'].font.size=Pt(8.5);styles['Caption'].font.color.rgb=RGBColor.from_string('55635C')
for name in ['List Number','List Bullet']:
 styles[name].paragraph_format.line_spacing=1.08;styles[name].paragraph_format.space_after=Pt(5)
h=S.header.paragraphs[0];h.text='NXT RETRO MUSIC  /  HƯỚNG DẪN v1.2';h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for run in h.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('6C7772')
f=S.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('NextUI & Stock   •   ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
for run in f.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('6C7772')
S.different_first_page_header_footer=True
D.core_properties.title=M['title']+' — '+M['subtitle'];D.core_properties.author='Nguyễn Xuân Thông'
D.core_properties.subject='Cài đặt, menu và phím điều khiển NXT Retro Music v1.2 trên TrimUI Brick Pro'
D.core_properties.language='vi-VN'
md=['# '+M['title'],'',M['subtitle'],'',M['edition'],'']
def table(block,page_index):
 spec=block['table'];T=D.add_table(rows=1,cols=len(spec['headers']));T.alignment=WD_TABLE_ALIGNMENT.CENTER;T.autofit=False
 for c,width in zip(T.columns,spec['widths']):c.width=Inches(width)
 T.style='Table Grid'
 for j,title in enumerate(spec['headers']):T.rows[0].cells[j].text=title
 for row in spec['rows']:
  cells=T.add_row().cells
  for j,value in enumerate(row):cells[j].text=value
 for i,row in enumerate(T.rows):
  no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
  for j,cell in enumerate(row.cells):
   cell.width=Inches(spec['widths'][j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr()
   margins=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:w'),'60' if side in ['top','bottom'] else '95');x.set(qn('w:type'),'dxa');margins.append(x)
   pr.append(margins)
   if i==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E2EEE7');pr.append(shade)
   elif i%2==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F6F8F6');pr.append(shade)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.04
    for run in p.runs:run.font.size=Pt(9 if page_index==19 else 9.35);run.font.bold=i==0
  if i==0:
   repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
 D.add_paragraph().paragraph_format.space_after=Pt(0)
 md.append('| '+' | '.join(spec['headers'])+' |');md.append('| '+' | '.join(['---']*len(spec['headers']))+' |')
 for row in spec['rows']:md.append('| '+' | '.join(v.replace('|','\\|').replace('\n','<br>') for v in row)+' |')
 md.append('')
for i,page in enumerate(M['pages']):
 if i:D.add_page_break()
 D.add_heading(page['title'],0 if page.get('cover') else 1)
 md+=['## '+page['title'],'']
 for b in page['blocks']:
  if 'p' in b:
   p=D.add_paragraph(b['p']);md+=[b['p'],'']
   if i==0 and page['blocks'].index(b)<2:
    for run in p.runs:run.font.size=Pt(14 if page['blocks'].index(b)==0 else 9.5)
  elif 'h' in b:D.add_heading(b['h'],2);md+=['### '+b['h'],'']
  elif 'table' in b:table(b,i)
  elif 'steps' in b:
   for j,value in enumerate(b['steps']):
    # Explicit numbering restarts at 1 within each procedure.
    p=D.add_paragraph();p.paragraph_format.left_indent=Inches(.2);p.paragraph_format.first_line_indent=Inches(-.2)
    p.add_run(str(j+1)+'. ').bold=True;p.add_run(value);md.append(str(j+1)+'. '+value)
   md.append('')
  elif 'image' in b:
   p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(R/b['image']),width=Inches(b['width']))
   cap=D.add_paragraph(b['caption'],'Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;md+=[b['caption'],'']
name='HUONG_DAN_CHI_TIET_NXT_Retro_Music_v1.2'
D.save(O/(name+'.docx'))
text='\n'.join(md)+'\n';(O/(name+'.md')).write_text(text)
(R/'HUONG_DAN_NXT_Retro_Music_v1.2_Stock.md').write_text(text)
print(O/(name+'.docx'))
print('Created',len(M['pages']),'explicit manual pages and matching Markdown.')
